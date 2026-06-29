"""Shared tabular-file reading (xlsx / csv) for catalog uploads.

Both the JD Union catalog check and the Taobao shop sync accept the same kind
of user-supplied product file, so the raw reading lives here once. Callers
apply their own column/URL extraction on top of the returned row matrix.
"""

from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Any


def read_xlsx(data: bytes) -> list[tuple[Any, ...]]:
    import openpyxl  # local import: heavy, only needed for xlsx

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    return list(ws.iter_rows(values_only=True))


def read_csv(text: str) -> list[tuple[Any, ...]]:
    reader = csv.reader(io.StringIO(text))
    return [tuple(r) for r in reader]


def read_tabular(data: bytes, filename: str) -> list[tuple[Any, ...]]:
    """Read xlsx or csv bytes into a header+data row matrix."""
    if filename.lower().endswith((".xlsx", ".xlsm")):
        return read_xlsx(data)
    return read_csv(data.decode("utf-8"))


def read_tabular_path(path: str | Path) -> list[tuple[Any, ...]]:
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        return read_xlsx(path.read_bytes())
    return read_csv(path.read_text("utf-8"))
